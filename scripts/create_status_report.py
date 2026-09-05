from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Users\Man$i verma\OneDrive\Documents\Gravity\Covenant_Certificate_Project_Status_Report.docx"
NAVY = "17365D"
PALE_BLUE = "EAF1F8"
PALE_GRAY = "F3F4F6"
BORDER = "D9D9D9"
TEXT = RGBColor(31, 41, 55)


def set_cell_fill(cell, color):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_borders(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_table(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[col_idx])
            if row_idx == 0:
                set_cell_fill(cell, NAVY)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9)
            else:
                set_cell_fill(cell, PALE_BLUE if row_idx % 2 == 0 else "FFFFFF")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        run.font.color.rgb = TEXT


def add_table(document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
    format_table(table, widths)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)
        paragraph.paragraph_format.space_after = Pt(3)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10)
normal.font.color.rgb = TEXT
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style.font.size = Pt(25)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(7)

for style_name, size in (("Heading 1", 16), ("Heading 2", 12)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.keep_with_next = True

for list_style in ("List Bullet", "List Number"):
    styles[list_style].font.name = "Aptos"
    styles[list_style].font.size = Pt(10)

title = doc.add_paragraph(style="Title")
title.add_run("Covenant Certificate Project Status Report")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = subtitle.add_run("Status as of 5 September 2026")
run.bold = True
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(75, 85, 99)

intro = doc.add_paragraph()
intro.add_run("Current conclusion. ").bold = True
intro.add_run(
    "The project has a working, tested hackathon prototype and the first production durability controls. "
    "The overall prototype is approximately 50 percent complete. Production readiness is approximately "
    "35 to 40 percent because the live persistence, upload, worker, durable review, event replay, and Gemini "
    "paths are not yet complete end to end."
)

doc.add_heading("Verified status", level=1)
add_table(
    doc,
    ["Area", "Verified result", "Status"],
    [
        ["Repository", "main at cbe2599 and synchronized with origin/main", "Passed"],
        ["Backend dependencies", "uv sync --frozen completed", "Passed"],
        ["Backend suite", "47 tests executed: 46 passed and 1 skipped", "Passed with one live-test gap"],
        ["Frontend dependencies", "npm ci completed", "Passed"],
        ["Frontend production build", "Next.js production build completed", "Passed"],
        ["Real Postgres restart test", "Test exists but TEST_DATABASE_URL was unavailable", "Not yet verified"],
        ["Supabase linked migration and lint checks", "CLI project link was unavailable on this machine", "Not rerun"],
        ["Real Gemini through LiteLLM", "No configured Gemini key was available", "Not tested"],
    ],
    [1.65, 4.15, 1.55],
)

doc.add_heading("Implemented work", level=1)
doc.add_heading("Covenant calculation and controls", level=2)
add_bullets(
    doc,
    [
        "Python FastAPI backend and LangGraph orchestration are in place.",
        "Deterministic Python performs covenant arithmetic with evidence checks; the model does not decide the final calculation.",
        "Missing inputs, unclear document precedence, unsupported adjustments, and incomplete review evidence fail closed instead of becoming zero or a pass.",
        "The workflow can return a reviewed draft pass, draft breach, or needs-review result without claiming full agreement compliance.",
        "Agreement-specific thresholds, amendments, citations, calculation lines, certificate draft data, and a hash-based audit trace are represented.",
    ],
)

doc.add_heading("Extraction and review workflow", level=2)
add_bullets(
    doc,
    [
        "A narrow Aon agreement PDF and financial extraction flow works for the curated fixture.",
        "Case revisions, change impacts, stale-command conflicts, review resolutions, snapshots, and officer approval logic exist in memory.",
        "Officer approval is bound to the exact revision, package hash, threshold, ratio, comparator, and financial inputs.",
        "A later revision invalidates the previous approval rather than silently reusing it.",
    ],
)

doc.add_heading("Supabase and durability foundation", level=2)
add_bullets(
    doc,
    [
        "Four forward migrations define organizations, memberships, cases, immutable documents, revisions, impacts, rules, facts, reviews, approvals, artifacts, events, a job queue, row-level security, and private Storage policies.",
        "The queue implementation includes leases, retry counts, heartbeat handling, cancellation, and stale-worker fencing tokens.",
        "Postgres and LangGraph checkpoint dependencies are pinned in the Python lockfile.",
        "Configured queue or checkpoint failures no longer silently fall back to memory.",
        "Liveness and dependency-aware readiness endpoints report queue and checkpoint status separately and redact connection details.",
        "LangGraph checkpoint schema setup was removed from application startup and remains an explicit operator action.",
    ],
)

doc.add_heading("Frontend and agent interface", level=2)
add_bullets(
    doc,
    [
        "The Next.js interface presents cases, formulas, financial evidence, cited clauses, results, blockers, reviewer actions, and workflow trace data.",
        "CopilotKit and AG-UI connect the frontend to the LangGraph agent endpoint.",
        "The Turbopack workspace root is fixed to the web application, making the local production build deterministic.",
        "Unused starter travel, todo, and chart examples were removed so the repository reflects the covenant product.",
    ],
)

doc.add_heading("Test report", level=1)
add_table(
    doc,
    ["Test area", "Coverage", "Result"],
    [
        ["Covenant core", "Opposite agreement outcomes, amendments, evidence gaps, unsupported covenants, safe empty scope, audit hashes", "Passed"],
        ["Review and approval", "Reviewer identity and rationale, stale decisions, idempotency, exact approval values, supersession", "Passed"],
        ["API and AG UI", "HTTP workflow, missing cases, health, AG-UI route registration", "Passed"],
        ["Platform adapters", "Offline mode, authentication requirement, safe observability behavior", "Passed"],
        ["Durability readiness", "Explicit offline mode, redacted failure, liveness, blocked business traffic", "Passed"],
        ["Postgres durability", "Queue and checkpoint reconstruction test", "Skipped pending isolated database"],
        ["Frontend", "Clean install, TypeScript, optimized Next.js build and routes", "Passed"],
    ],
    [1.55, 4.45, 1.35],
)

doc.add_heading("Completion estimate", level=1)
add_table(
    doc,
    ["Workstream", "Estimated completion", "Assessment"],
    [
        ["Deterministic covenant core", "85 percent", "Working and well covered for curated leverage cases"],
        ["Database schema and security foundation", "80 percent", "Schema and policies exist; full runtime use remains"],
        ["Durability and readiness foundation", "70 percent", "Fail-closed behavior works; live restart proof remains"],
        ["Revision, review, and approval", "55 percent", "Strong in-memory behavior; Postgres repository replacement remains"],
        ["Document intake and Storage", "30 percent", "Narrow parser exists; authenticated immutable upload is missing"],
        ["Production worker", "30 percent", "Queue mechanics exist; worker execution is not connected end to end"],
        ["Durable human interrupt and resume", "20 percent", "Review logic exists, but paused graph recovery is missing"],
        ["Durable AG UI events and replay", "20 percent", "Basic transport exists; outbox, sequence, snapshot replay remain"],
        ["Complete production frontend", "35 percent", "Demo dashboard works; upload-through-approval flow remains"],
        ["Gemini integration and accuracy evaluation", "15 percent", "Aliases and labels exist; live calls and full harness remain"],
    ],
    [2.45, 1.5, 3.4],
)

summary = doc.add_paragraph()
summary.add_run("Overall engineering estimate. ").bold = True
summary.add_run(
    "Approximately 50 percent of the hackathon prototype is implemented. Around 50 percent remains for a complete "
    "demonstration workflow. For production readiness, approximately 60 to 65 percent remains because live infrastructure "
    "and recovery behavior carry more risk than the existing interface and deterministic calculation code."
)

doc.add_heading("Remaining work", level=1)
remaining = [
    "Replace the global in-memory revision store with organization-scoped Supabase Postgres repositories.",
    "Authenticate every private read and mutation and derive actor, organization, and role from verified server identity.",
    "Implement authenticated multipart upload, content validation, SHA-256 streaming, immutable private Storage, document versions, revisions, events, and queued jobs in one controlled flow.",
    "Run a real worker that leases jobs, heartbeats, retries transient failures, recovers after restart, and prevents stale workers from publishing.",
    "Persist extracted rules and facts as authoritative revision state and recalculate only after deterministic validation.",
    "Add durable LangGraph interrupt and resume for matching human review decisions.",
    "Persist redacted domain events and add ordered AG-UI snapshot and replay after browser reconnect.",
    "Complete the frontend from Supabase sign-in and case creation through upload, review, recalculation, and exact officer approval.",
    "Run authenticated Gemini calls through LiteLLM and verify tool use, structured output, retries, provider reporting, latency, tokens, and cost.",
    "Complete the accuracy harness, holdout agreement family, citation checks, false-pass gates, security checks, PDF workpaper export, staging deployment, and restart runbook.",
]
for idx, item in enumerate(remaining, 1):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.add_run(item)
    paragraph.paragraph_format.space_after = Pt(3)

doc.add_heading("Immediate priorities", level=1)
add_table(
    doc,
    ["Priority", "Next deliverable", "Acceptance evidence"],
    [
        ["1", "Postgres revision repositories and verified identity", "Restart persistence, tenant isolation, role enforcement, stale 409 behavior"],
        ["2", "Immutable upload and worker connection", "One upload creates object, version, revision, event, job, and recoverable result"],
        ["3", "Durable review and event replay", "Paused run survives restart; reconnect replays each event once in order"],
        ["4", "Complete frontend workflow", "Reviewer completes upload through current-revision approval on desktop and mobile"],
        ["5", "Gemini and evaluation proof", "Real LiteLLM call tests plus versioned accuracy metrics and zero observed unsupported passes"],
    ],
    [0.75, 2.8, 3.8],
)

doc.add_heading("Current constraints", level=1)
constraints = doc.add_paragraph()
constraints.add_run("Live verification remains limited. ").bold = True
constraints.add_run(
    "The latest local run did not have an isolated Postgres test database, a linked Supabase CLI session, or a configured "
    "Gemini key. The existing active tests and frontend build are green, but the report does not claim that live database "
    "restart recovery, hosted row-level security, Storage isolation, or real model calls have passed. Server credentials "
    "must remain in ignored environment files or a secret manager and must never be committed."
)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Covenant Certificate Project Status  |  5 September 2026")
footer_run.font.name = "Aptos"
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(107, 114, 128)

doc.core_properties.title = "Covenant Certificate Project Status Report"
doc.core_properties.subject = "Implementation progress test results and remaining work"
doc.core_properties.author = "Covenant Certificate Team"
doc.save(OUTPUT)
print(OUTPUT)
